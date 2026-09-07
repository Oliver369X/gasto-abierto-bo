"""Analyze auditoria-drive-MASTER workbook for field/KPI/join requirements.

Usage:
  python scripts/analyze_audit_workbook.py
  python scripts/analyze_audit_workbook.py path/to/workbook.xlsx

Place Diego's file at:
  docs/requirements/auditoria-drive-MASTER-v9.3.xlsx

Requires: pip install openpyxl  (or pip install -e ".[workbook]")
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_PATH = ROOT / "docs" / "requirements" / "auditoria-drive-MASTER-v9.3.xlsx"

# Map header tokens → Gasto Abierto schema / API (requirements traceability)
FIELD_HINTS: dict[str, tuple[str, str]] = {
    r"entidad|institucion|organismo|ministerio|gobernacion|alcaldia": (
        "entity.name",
        "Institución — join Presupuesto Abierto + SICOES",
    ),
    r"gestion|anio|ejercicio|year": ("budget_line.year", "Gestión fiscal"),
    r"departamento|ubicacion|geograf|municipio|territorio": (
        "entity.department / territory",
        "Geografía — clasificador territorial MEFP",
    ),
    r"objeto|partida|clasificador|rubro|programa|proyecto|finalidad|sector": (
        "budget_line.program_project / budget_item / category",
        "Objeto de gasto / programa",
    ),
    r"cuce|proceso|contrato|adjudic": ("contract.cuce", "Cruce CUCE SICOES ↔ OCP"),
    r"proveedor|adjudicatario|contratista|nit": (
        "supplier.name / nit",
        "Proveedor — supplier_master",
    ),
    r"presupuesto|vigente|inicial|modificad|ppto": (
        "budget_line.current_amount / initial_amount",
        "Fases presupuesto (no sumar entre sí)",
    ),
    r"ejecuc|deveng|pagad|comprom": (
        "budget_line.payment / executed_amount / commitment",
        "Ejecución presupuestaria",
    ),
    r"discrep|diferencia|delta|variacion": (
        "discrepancy",
        "Cruce fuentes — /v1/discrepancies",
    ),
    r"alerta|riesgo|anomalia|concentracion|directa": (
        "alert.rule_id",
        "Reglas ciudadanas — /v1/alerts",
    ),
    r"auditor|hallazgo|contralor|informe": (
        "audit_report / audit_finding",
        "Contraloría — /v1/audits",
    ),
    r"fuente|origen|sigep|sicoes|presupuesto.?abierto": (
        "source_id / ingestion_run_id",
        "Provenance — docs/sources/",
    ),
}

KPI_PATTERNS: dict[str, str] = {
    r"ratio|porcent|%|ejecucion": "execution_ratio_pct — /v1/budgets/totals",
    r"total|suma|monto|importe": "aggregate totals — /v1/budgets/aggregate",
    r"compar|histor|serie|tendencia|yoy": "history — /v1/history/years",
    r"cobertura|complet": "coverage — /v1/coverage/sicoes, /v1/product-gate",
}


def _norm(s: str) -> str:
    return (
        s.strip()
        .lower()
        .replace("á", "a")
        .replace("é", "e")
        .replace("í", "i")
        .replace("ó", "o")
        .replace("ú", "u")
        .replace("ñ", "n")
    )


def classify_header(header: str) -> list[dict[str, str]]:
    h = _norm(header)
    out: list[dict[str, str]] = []
    for pattern, (field, note) in FIELD_HINTS.items():
        if re.search(pattern, h):
            out.append({"header": header, "maps_to": field, "note": note})
            break
    for pattern, api in KPI_PATTERNS.items():
        if re.search(pattern, h):
            out.append({"header": header, "kpi_hint": api})
            break
    return out


def analyze(path: Path) -> dict:
    try:
        import openpyxl  # type: ignore[import-untyped]
    except ImportError as exc:
        raise SystemExit(
            "Instalá openpyxl: pip install openpyxl  (o pip install -e '.[workbook]')"
        ) from exc

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets_out: list[dict] = []
    all_mappings: list[dict] = []
    suggested_joins: set[str] = set()

    for name in wb.sheetnames:
        ws = wb[name]
        rows = ws.iter_rows(max_row=5, values_only=True)
        header_row = next(rows, None)
        headers = [str(c).strip() for c in (header_row or []) if c is not None and str(c).strip()]
        mappings = []
        for h in headers:
            for m in classify_header(h):
                mappings.append(m)
                if "maps_to" in m:
                    all_mappings.append({**m, "sheet": name})
                    if "entity" in m["maps_to"] and "year" in str(m.get("maps_to", "")):
                        suggested_joins.add("entity + year")
                    if "cuce" in m["maps_to"]:
                        suggested_joins.add("cuce cross-source")
                    if "department" in m["maps_to"]:
                        suggested_joins.add("institución + geografía")
        sample_rows = 0
        for _ in ws.iter_rows(min_row=2, max_row=6, values_only=True):
            sample_rows += 1
        sheets_out.append(
            {
                "sheet": name,
                "headers": headers,
                "header_count": len(headers),
                "sample_data_rows": sample_rows,
                "mappings": mappings,
            }
        )
    wb.close()

    # Infer joins from co-occurring dimensions on same sheet
    for sh in sheets_out:
        fields = {m.get("maps_to", "") for m in sh["mappings"] if "maps_to" in m}
        txt = " ".join(fields)
        if "entity.name" in txt and "budget_line.year" in txt:
            suggested_joins.add(f"{sh['sheet']}: entidad × gestión")
        if "entity.name" in txt and "contract.cuce" in txt:
            suggested_joins.add(f"{sh['sheet']}: entidad × contrato (CUCE)")
        if "department" in txt and "budget_line" in txt:
            suggested_joins.add(f"{sh['sheet']}: geografía × presupuesto")

    official_sources = [
        "presupuesto_abierto (CSV/Parquet abierto.economiayfinanzas.gob.bo/descargas)",
        "sicoes / agetic (datos.gob.bo, OCP)",
        "cge (Contraloría metadatos públicos)",
    ]

    return {
        "workbook": str(path),
        "sheet_count": len(sheets_out),
        "sheets": sheets_out,
        "field_mappings": all_mappings,
        "suggested_joins": sorted(suggested_joins),
        "official_data_priority": official_sources,
        "platform_coverage": {
            "presupuesto_aggregate": "/v1/budgets/aggregate",
            "history_by_entity": "/v1/history/years?entity_id=",
            "cross_source": "/v1/cross-source/summary",
            "discrepancies": "/v1/discrepancies",
        },
    }


def main() -> int:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_PATH
    if not path.exists():
        print(
            json.dumps(
                {
                    "error": "workbook_not_found",
                    "expected_path": str(DEFAULT_PATH),
                    "hint": "Copiá auditoria-drive-MASTER-v9.3.xlsx a docs/requirements/ y re-ejecutá.",
                },
                indent=2,
                ensure_ascii=False,
            )
        )
        return 2
    result = analyze(path)
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
